"""One Excel workbook with everything (needs openpyxl): summary, libraries, recipes, nutrition."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from urllib.parse import quote

from .db import DB

FONT = "Arial"


def export_xlsx(db: DB, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    sources = [r[0] for r in db.query("SELECT DISTINCT source FROM recipes ORDER BY source")]
    wb = Workbook()
    head_font = Font(name=FONT, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3B5B4F")
    body_font = Font(name=FONT, size=10)

    link_font = Font(name=FONT, size=10, color="0563C1", underline="single")

    def sheet(title: str, header: list[str], rows, widths: dict[int, int] | None = None, first: bool = False,
              links: tuple[str, ...] = ()):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for row in rows:
            ws.append(["" if v is None else v for v in row])
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        link_cols = {i for i, h in enumerate(header) if h in links}
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.column - 1 in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
                    # tap to open the photo / the page it was published on; the target is percent-encoded
                    # because some phone apps won't open links containing Georgian letters
                    cell.hyperlink = quote(cell.value, safe=":/?=&%#+,;@~")
                    cell.font = link_font
                else:
                    cell.font = body_font
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for i, _ in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)
        return ws

    def j(value: str | None, sep: str = " | ") -> str:
        if not value:
            return ""
        data = json.loads(value)
        if isinstance(data, dict):
            return sep.join(f"{k}: {v}" for k, v in data.items())
        return sep.join(str(x) for x in data)

    # ---- Summary. Plain values, not formulas: openpyxl can't store computed results, and formula cells
    # without them show up blank in phone previewers until Excel recalculates.
    ws = sheet("Summary", ["site", "recipes", "ingredient lines", "dishes on this site", "ingredients on this site",
                           "failed pages"], [], {1: 16, 2: 12, 3: 16, 4: 18, 5: 22, 6: 13}, first=True)
    n_ing = db.query("SELECT COUNT(*) FROM ingredients")[0][0]
    n_dish = db.query("SELECT COUNT(*) FROM dishes")[0][0]
    per_site = {s: [0, 0, 0, 0, 0] for s in sources}
    for r in db.query("SELECT source, COUNT(*), COALESCE(SUM(n_ingredients), 0) FROM recipes GROUP BY source"):
        per_site[r[0]][0:2] = [r[1], r[2]]
    for table, col in (("dishes", 2), ("ingredients", 3)):
        for r in db.query(f"SELECT per_site FROM {table}"):
            for s in json.loads(r[0] or "{}"):
                if s in per_site:
                    per_site[s][col] += 1
    for r in db.query("SELECT source, COUNT(*) FROM failures GROUP BY source"):
        if r[0] in per_site:
            per_site[r[0]][4] = r[1]
    for s in sources:
        ws.append([s, *per_site[s]])
    ws.append(["all sites", sum(v[0] for v in per_site.values()), sum(v[1] for v in per_site.values()),
               n_dish, n_ing, sum(v[4] for v in per_site.values())])
    total = len(sources) + 2
    for cell in ws[total]:
        cell.font = Font(name=FONT, size=10, bold=True)
    notes = [
        "",
        "How to read this workbook",
        "Ingredient library: one row per distinct ingredient across all sites (Georgian case forms and prep words merged).",
        "  in_<site> = number of that site's recipes using it. kcal/protein/fat/carbs are per 100 g from the gemrielia.ge "
        "and fiber.ge nutrition tables, matched by name — check before using.",
        "Dish library: one row per distinct dish (clickbait removed from titles); in_<site> = recipes on that site.",
        "Recipe ingredients: every ingredient line as published (raw) and parsed (name, quantity, unit, grams/ml).",
        "  ml uses 1 tbsp (ს/კ) = 15 ml, 1 tsp (ჩ/კ) = 5 ml, 1 cup (ჭიქა) = 250 ml. grams are only filled from weights.",
        "Recipes: calories_site is what the site prints (usually per serving), not computed.",
        "Photos: one row per recipe photo with the page it was published on (tap a link to open it); 'dish photo' = "
        "yes marks the photo representing the dish. Photos belong to the sites/authors — get permission before "
        "showing them in an app.",
        f"Counts as of {date.today().isoformat()}.",
    ]
    for line in notes:
        ws.append([line])
    ws.cell(row=total + 2, column=1).font = Font(name=FONT, bold=True)

    # ---- Ingredient library
    ing = db.query("SELECT * FROM ingredients ORDER BY recipes DESC, name")
    sheet("Ingredient library",
          ["name", "recipes", "sites", "key", *[f"in_{s}" for s in sources], "kcal_100g", "protein_100g", "fat_100g",
           "carbs_100g", "nutrition_match", "spelling variants", "units used", "examples"],
          ([r["name"], r["recipes"], r["sites"], r["key"], *[json.loads(r["per_site"]).get(s, 0) for s in sources],
            r["kcal_100g"], r["protein_100g"], r["fat_100g"], r["carbs_100g"],
            f"{r['nutrition_name']} ({r['nutrition_source']})" if r["nutrition_name"] else "",
            j(r["variants"]), j(r["units"]), j(r["examples"])] for r in ing),
          {1: 28, 4: 18, 9 + len(sources): 30, 10 + len(sources): 40, 11 + len(sources): 24, 12 + len(sources): 60})

    # ---- Dish library
    dishes = db.query("SELECT d.*, i.path AS image_file FROM dishes d LEFT JOIN images i "
                      "ON i.recipe_id = d.image_recipe_id ORDER BY d.sites DESC, d.recipes DESC, d.name")
    n = len(sources)
    sheet("Dish library",
          ["dish", "recipes", "sites", "key", *[f"in_{s}" for s in sources], "title variants", "categories",
           "top ingredients", "photo url", "photo file", "example url"],
          ([r["name"], r["recipes"], r["sites"], r["key"], *[json.loads(r["per_site"]).get(s, 0) for s in sources],
            j(r["variants"]), j(r["categories"]), j(r["top_ingredients"], ", "), r["image"], r["image_file"],
            (json.loads(r["urls"]) or [""])[0]] for r in dishes),
          {1: 34, 4: 18, 5 + n: 50, 6 + n: 30, 7 + n: 60, 8 + n: 50, 9 + n: 30, 10 + n: 50},
          links=("photo url", "example url"))

    # ---- Photos: every photo and the page it was published on
    from .export import PHOTO_HEADER, photo_rows
    sheet("Photos", [h.replace("_", " ") for h in PHOTO_HEADER], photo_rows(db),
          {1: 34, 2: 40, 3: 12, 4: 60, 5: 60, 6: 16, 7: 11, 8: 30}, links=("photo url", "published on"))

    # ---- Recipes
    recipes = db.query("SELECT r.*, i.path AS image_file FROM recipes r LEFT JOIN images i ON i.recipe_id = r.id "
                       "ORDER BY r.source, r.id")
    sheet("Recipes",
          ["id", "source", "title", "dish", "categories", "servings", "minutes", "calories_site", "ingredients", "url",
           "photo url", "photo file"],
          ([r["id"], r["source"], r["title"], r["dish_name"], j(r["categories"]), r["servings"], r["total_minutes"],
            r["calories"], r["n_ingredients"], r["url"], r["image"], r["image_file"]] for r in recipes),
          {1: 7, 3: 40, 4: 30, 5: 30, 10: 60, 11: 60, 12: 30}, links=("url", "photo url"))

    # ---- Recipe ingredients
    titles = {r["id"]: r["title"] for r in recipes}
    source_of = {r["id"]: r["source"] for r in recipes}
    lines = db.query("SELECT ri.*, i.name AS cname FROM recipe_ingredients ri LEFT JOIN ingredients i "
                     "ON i.key = ri.canonical ORDER BY ri.recipe_id, ri.position")
    sheet("Recipe ingredients",
          ["recipe_id", "source", "recipe", "raw", "name", "library name", "quantity", "quantity_max", "unit",
           "unit_norm", "grams", "ml", "note", "group", "optional"],
          ([r["recipe_id"], source_of.get(r["recipe_id"]), titles.get(r["recipe_id"]), r["raw"], r["name"], r["cname"],
            r["quantity"], r["quantity_max"], r["unit"], r["unit_norm"], r["grams"], r["ml"], r["note"], r["grp"],
            "yes" if r["optional"] else ""] for r in lines),
          {3: 34, 4: 40, 5: 26, 6: 26, 13: 26})

    # ---- Nutrition tables
    nut = db.query("SELECT * FROM nutrition ORDER BY source, grp, name")
    if nut:
        sheet("Nutrition tables", ["source", "group", "name", "kcal_100g", "protein_100g", "fat_100g", "carbs_100g"],
              ([r["source"], r["grp"], r["name"], r["kcal"], r["protein"], r["fat"], r["carbs"]] for r in nut),
              {2: 40, 3: 36})

    wb.save(path)
    return path
